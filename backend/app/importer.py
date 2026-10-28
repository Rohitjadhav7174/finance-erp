"""Excel import: parse → preview (validation, duplicates, reconciliation) → confirm → DRAFT entries.

Safety rules
  * Nothing is posted by an import. Confirmed rows become DRAFT entries that an
    accountant reviews and posts like any other entry.
  * Original cell values, cell addresses, number formats and comments are stored
    unchanged in fin.import_rows.raw (the database refuses to modify them).
  * Missing information is never invented: a row that lacks a required choice
    (customer, cash/bank account, …) is BLOCKED until a person supplies it.
  * Workbook totals (formula cells) are compared with what will be imported.

Profiles
  * sindhudurg_balance_sheet — the owner's "Monthly Balance Sheet" layout. Decisions
    recorded in docs/01 §6: column C = money RECEIVED from a customer (I-1); rows with a
    category in D and an amount in E = EXPENSES PAID by cash or bank (I-2).
  * ledger — one row per transaction (Date, Type, Description, Party, Bill Raised,
    Installment Received, Expense Paid, Payment Mode, Invoice/Receipt No., Notes).
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

MAX_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED = 60 * 1024 * 1024
MAX_ROWS = 5000


class ImportFileError(ValueError):
    pass


@dataclass
class ParsedRow:
    sheet: str
    row_number: int
    raw: dict[str, Any]
    mapped: dict[str, Any] | None  # None => informational row (totals, headings)
    issues: list[dict] = field(default_factory=list)


@dataclass
class ControlTotal:
    sheet: str
    cell: str
    formula: str
    excel_value: Decimal | None
    recomputed: Decimal | None
    covers_column: str | None
    covers_rows: list[int]
    note: str | None = None


@dataclass
class Parsed:
    profile: str
    sha256: str
    rows: list[ParsedRow]
    controls: list[ControlTotal]
    sheets: list[dict]


def issue(level: str, code: str, message: str) -> dict:
    return {"level": level, "code": code, "message": message}


# ---------------------------------------------------------------------------
# File checks
# ---------------------------------------------------------------------------
def check_file(name: str, data: bytes) -> str:
    if not name.lower().endswith(".xlsx"):
        raise ImportFileError("Only .xlsx Excel files can be imported (not .xls, .xlsm or .csv).")
    if len(data) > MAX_BYTES:
        raise ImportFileError("The file is larger than 5 MB.")
    if not data.startswith(b"PK\x03\x04"):
        raise ImportFileError("This is not a valid .xlsx file.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
                raise ImportFileError("The file is too large when unpacked.")
            names = {i.filename for i in infos}
            if "xl/vbaProject.bin" in names:
                raise ImportFileError("Files containing macros are not accepted.")
            if "xl/workbook.xml" not in names:
                raise ImportFileError("This is not a valid .xlsx file.")
    except zipfile.BadZipFile:
        raise ImportFileError("This is not a valid .xlsx file.") from None
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------
def _raw_cell(c) -> dict:
    v = c.value
    if isinstance(v, datetime):
        v = v.isoformat()
    elif isinstance(v, date):
        v = v.isoformat()
    elif isinstance(v, Decimal):
        v = str(v)
    d = {"v": v, "fmt": c.number_format}
    if c.comment:
        d["comment"] = c.comment.text
    return d


def _money(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise InvalidOperation
    if isinstance(v, (int, float)):
        d = Decimal(str(v))
    else:
        d = Decimal(str(v).replace(",", "").replace("₹", "").strip())
    if d != d.quantize(Decimal("0.01")):
        raise InvalidOperation
    return d.quantize(Decimal("0.01"))


def _date(v, fmt: str, issues: list[dict], label: str) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        d = from_excel(v).date()
        issues.append(issue("warning", "serial_date",
                            f"{label} was stored as the number {v:g} (format '{fmt}'); read as the date {d:%d-%m-%Y}. Please check."))
        return d
    if isinstance(v, str):
        for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    issues.append(issue("error", "bad_date", f"{label} '{v}' is not a date."))
    return None


CATEGORY_CODES = [("godown", "EXP-GDN"), ("diesel", "EXP-DSL"), ("employee", "EXP-EMP"), ("worker", "EXP-WRK")]
METHODS = [("cash", "cash"), ("upi", "upi"), ("gpay", "upi"), ("phonepe", "upi"), ("neft", "bank_transfer"),
           ("rtgs", "bank_transfer"), ("imps", "bank_transfer"), ("bank", "bank_transfer"), ("cheque", "cheque"),
           ("check", "cheque")]


def _clean_comment(text: str) -> str:
    """Readable note text; the untouched original stays in import_rows.raw."""
    keep = [ln.strip() for ln in text.splitlines()
            if ln.strip() and not re.match(r"^(=+|ID#\S+|Unknown Author\b.*|-+)$", ln.strip())]
    return " ".join(keep)


def guess_category(text: str | None) -> str | None:
    t = (text or "").lower()
    return next((code for key, code in CATEGORY_CODES if key in t), None)


def guess_method(text: str | None) -> str | None:
    t = (text or "").lower()
    return next((m for key, m in METHODS if key in t), None)


# ---------------------------------------------------------------------------
# Formula control totals (SUM ranges and A+B+C) — recomputed from raw values
# ---------------------------------------------------------------------------
SUM_RE = re.compile(r"^=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$", re.I)
PLUS_RE = re.compile(r"^=([A-Z]+\d+)(\+[A-Z]+\d+)+$", re.I)
REF_RE = re.compile(r"([A-Z]+)(\d+)", re.I)


def _controls(ws_f, ws_v) -> list[ControlTotal]:
    out = []
    for row in ws_f.iter_rows():
        for c in row:
            f = c.value
            if not (isinstance(f, str) and f.startswith("=")):
                continue
            cached = ws_v[c.coordinate].value
            excel_value = _try_money(cached)
            refs: list[tuple[str, int]] = []
            m = SUM_RE.match(f.replace(" ", ""))
            if m and m.group(1).upper() == m.group(3).upper():
                refs = [(m.group(1).upper(), r) for r in range(int(m.group(2)), int(m.group(4)) + 1)]
            elif PLUS_RE.match(f.replace(" ", "")):
                refs = [(a.upper(), int(b)) for a, b in REF_RE.findall(f)]
            if not refs or len({col for col, _ in refs}) != 1:
                continue  # only simple column totals are reconciled
            col = refs[0][0]
            total = Decimal(0)
            for _, r in refs:
                v = _try_money(ws_v[f"{col}{r}"].value)
                total += v or 0
            # rows with amounts in this column above the total that the formula does not include (I-9);
            # filtered against the other totals of the column after the loop
            rows_in = {r for _, r in refs}
            left_out = [r for r in range(3, c.row) if r not in rows_in and _try_money(ws_v[f"{col}{r}"].value)
                        and not (isinstance(ws_f[f"{col}{r}"].value, str) and ws_f[f"{col}{r}"].value.startswith("="))]
            note = ",".join(map(str, left_out)) or None
            out.append(ControlTotal(ws_f.title, c.coordinate, f, excel_value, total.quantize(Decimal("0.01")),
                                    col, sorted(rows_in), note))
    for ct in out:
        if ct.note:
            covered = {r for other in out if other is not ct and other.covers_column == ct.covers_column
                       for r in other.covers_rows}
            missing = [int(r) for r in ct.note.split(",") if int(r) not in covered]
            ct.note = (f"This total leaves out rows {', '.join(map(str, missing))} that have amounts in column "
                       f"{ct.covers_column}." if missing else None)
    return out


def _try_money(v) -> Decimal | None:
    try:
        return _money(v) if isinstance(v, (int, float, Decimal)) or (isinstance(v, str) and v.strip()) else None
    except (InvalidOperation, ValueError):
        return None


# ---------------------------------------------------------------------------
# Profiles
# ---------------------------------------------------------------------------
def _header(ws, row=2) -> list[str]:
    return [str(c.value or "").strip().lower() for c in ws[row]]


def _is_balance_sheet(ws) -> bool:
    h = _header(ws)
    return len(h) >= 5 and h[0].startswith("period") and "opening balance" in h[2] and h[3].startswith("payment")


def _ledger_columns(ws) -> dict[str, int] | None:
    h = _header(ws)
    want = {"date": "date", "type": "type", "description": "description", "party": "party", "bill": "bill raised",
            "installment": "installment received", "expense": "expense paid", "mode": "payment mode",
            "ref": "invoice / receipt", "notes": "notes"}
    cols = {}
    for key, needle in want.items():
        idx = next((i for i, t in enumerate(h) if t.startswith(needle)), None)
        if idx is not None:
            cols[key] = idx
    return cols if {"date", "type"} <= cols.keys() and ({"installment", "expense", "bill"} & cols.keys()) else None


def parse(name: str, data: bytes, today: date | None = None) -> Parsed:
    sha = check_file(name, data)
    today = today or date.today()
    try:
        wb_f = load_workbook(io.BytesIO(data), data_only=False)
        wb_v = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise ImportFileError("The workbook could not be opened. Is it a normal Excel file?") from None
    rows: list[ParsedRow] = []
    controls: list[ControlTotal] = []
    sheets = []
    profiles = set()
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        if ws_f.max_row > MAX_ROWS:
            raise ImportFileError(f"Sheet '{ws_f.title}' has more than {MAX_ROWS} rows.")
        title = str(ws_f["A1"].value or "").strip()
        if _is_balance_sheet(ws_f):
            profiles.add("sindhudurg_balance_sheet")
            sheet_rows = _parse_balance_sheet(ws_f, ws_v, today)
        elif (cols := _ledger_columns(ws_f)) is not None:
            profiles.add("ledger")
            sheet_rows = _parse_ledger(ws_f, ws_v, cols, today)
        else:
            sheets.append({"sheet": ws_f.title, "title": title, "profile": None, "data_rows": 0,
                           "message": "Layout not recognised — nothing imported from this sheet."})
            continue
        data_rows = sum(1 for r in sheet_rows if r.mapped)
        sheets.append({"sheet": ws_f.title, "title": title, "profile": "sindhudurg_balance_sheet"
                       if _is_balance_sheet(ws_f) else "ledger", "data_rows": data_rows,
                       "message": None if data_rows else "No transactions found on this sheet (template or empty)."})
        rows += sheet_rows
        controls += _controls(ws_f, ws_v)
    if not sheets:
        raise ImportFileError("The workbook has no sheets.")
    return Parsed("+".join(sorted(profiles)) or "unknown", sha, rows, controls, sheets)


def _common_date_checks(d: date | None, period: date | None, today: date, issues: list[dict]) -> None:
    if d is None:
        return
    if d > today:
        issues.append(issue("error", "future_date", f"Date {d:%d-%m-%Y} is in the future. Correct it before importing."))
    if period and (d.year, d.month) != (period.year, period.month):
        issues.append(issue("warning", "period_mismatch",
                            f"Date {d:%d-%m-%Y} is not in the row's period {period:%b-%Y}."))


def _parse_balance_sheet(ws_f, ws_v, today: date) -> list[ParsedRow]:
    out: list[ParsedRow] = []
    period: date | None = None
    for r in range(3, ws_f.max_row + 1):
        cells = {f"{col}{r}": _raw_cell(ws_f[f"{col}{r}"]) for col in "ABCDEFGHIJKLMN"
                 if ws_f[f"{col}{r}"].value is not None or ws_f[f"{col}{r}"].comment}
        if not cells:
            continue
        a = ws_v[f"A{r}"].value
        if isinstance(a, str):
            label = a.strip()
            if label.lower().startswith("total") or label.lower() in ("company details", "company", "address", "bank"):
                out.append(ParsedRow(ws_f.title, r, cells, None, [issue("info", "not_imported",
                                     f"'{label}' row — used for checking only, not imported.")]))
                continue
        row_issues: list[dict] = []
        if a is not None:
            p = _date(a, ws_f[f"A{r}"].number_format, row_issues, "Period")
            if p:
                period = p.replace(day=1)
        b = ws_v[f"B{r}"].value
        entry_date = _date(b, ws_f[f"B{r}"].number_format, row_issues, "Date")
        has_expense = ws_v[f"E{r}"].value not in (None, "")

        def comments_for(columns: str) -> list[str]:
            return [f"Workbook comment on {k}: {_clean_comment(v['comment'])}" for k, v in cells.items()
                    if v.get("comment") and k[0] in columns]
        # a comment stays with the column it was written in; D/E comments go to the expense row if there is one
        receipt_comments = comments_for("ABCFGHIJKLMN" + ("" if has_expense else "DE"))
        expense_comments = comments_for("DE")
        comment_issue = issue("warning", "comment_refers_to_missing_bill",
                              "The workbook has a comment here about notebook bill figures that are not in the sheet "
                              "(question I-3). The comment is kept as a note on the entry.")

        # Column C: money received from a customer (decision I-1)
        try:
            c_amt = _money(ws_v[f"C{r}"].value)
        except (InvalidOperation, ValueError):
            c_amt = None
            row_issues.append(issue("error", "bad_amount", f"C{r} is not a valid amount."))
        if c_amt is not None:
            iss = list(row_issues) + ([comment_issue] if receipt_comments else [])
            if entry_date is None:
                iss.append(issue("error", "missing_date", "Date (column B) is empty."))
            _common_date_checks(entry_date, period, today, iss)
            if c_amt <= 0:
                iss.append(issue("error", "bad_amount", "Amount must be greater than zero."))
            out.append(ParsedRow(ws_f.title, r, cells, {
                "kind": "receipt", "receipt_kind": "customer_payment", "entry_date": entry_date and str(entry_date),
                "period": period and str(period), "amount": str(c_amt), "party_id": None, "money_account_id": None,
                "payment_method": None, "reference_no": None, "source_cell": f"C{r}",
                "description": f"Received from customer (imported {ws_f.title.strip()} row {r})",
                "notes": "\n".join(receipt_comments) or None,
            }, iss))

        # Columns D + E: expense paid, category from D (decision I-2)
        category = ws_v[f"D{r}"].value
        e_val = ws_v[f"E{r}"].value
        if category and e_val in (None, ""):
            out.append(ParsedRow(ws_f.title, r, cells if c_amt is None else {k: v for k, v in cells.items() if k[0] in "DE"},
                                 None, [issue("info", "category_without_amount",
                                              f"'{category}' in D{r} has no amount in E{r} — nothing to import.")]))
        elif e_val not in (None, ""):
            iss = list(row_issues) + ([comment_issue] if expense_comments else [])
            try:
                e_amt = _money(e_val)
            except (InvalidOperation, ValueError):
                e_amt = None
                iss.append(issue("error", "bad_amount", f"E{r} is not a valid amount."))
            exp_date = entry_date
            if exp_date is None and period is not None:
                # Expense lines below the dated row inherit that row's date (rows 4-6 in the workbook).
                prev = next((x for x in reversed(out) if x.mapped and x.mapped.get("entry_date")), None)
                exp_date = date.fromisoformat(prev.mapped["entry_date"]) if prev else None
                if exp_date:
                    iss.append(issue("warning", "date_inherited",
                                     f"No date on this row; using {exp_date:%d-%m-%Y} from the row above. Please check."))
            if exp_date is None:
                iss.append(issue("error", "missing_date", "No date for this expense."))
            _common_date_checks(exp_date, period, today, iss)
            if not category:
                iss.append(issue("error", "missing_category", "Amount without a category in column D."))
            out.append(ParsedRow(ws_f.title, r, cells if c_amt is None else {k: v for k, v in cells.items() if k[0] in "DE"}, {
                "kind": "expense", "entry_date": exp_date and str(exp_date), "period": period and str(period),
                "amount": e_amt and str(e_amt), "category_text": category, "category_code": guess_category(category),
                "expense_account_id": None, "money_account_id": None, "payment_method": None, "party_id": None,
                "reference_no": None, "source_cell": f"E{r}",
                "description": f"{category} (imported {ws_f.title.strip()} row {r})",
                "notes": "\n".join(expense_comments) or None,
            }, iss))
    return out


def _parse_ledger(ws_f, ws_v, cols: dict[str, int], today: date) -> list[ParsedRow]:
    out: list[ParsedRow] = []
    letters = "ABCDEFGHIJKLMNOP"
    for r in range(3, ws_f.max_row + 1):
        row_v = [c.value for c in ws_v[r]]

        def get(k):
            i = cols.get(k)
            return row_v[i] if i is not None and i < len(row_v) else None
        cells = {f"{letters[i]}{r}": _raw_cell(ws_f[f"{letters[i]}{r}"]) for i in range(min(len(row_v), 16))
                 if ws_f[f"{letters[i]}{r}"].value is not None and not str(ws_f[f"{letters[i]}{r}"].value).startswith("=IF(")}
        typ = str(get("type") or "").strip().lower()
        amounts = {k: get(k) for k in ("bill", "installment", "expense") if get(k) not in (None, "")}
        if not typ and not amounts:
            if isinstance(get("date"), str) and get("date").strip().lower().startswith("total"):
                out.append(ParsedRow(ws_f.title, r, cells, None, [issue("info", "not_imported", "Total row — not imported.")]))
            continue
        iss: list[dict] = []
        d = _date(get("date"), ws_f.cell(r, cols["date"] + 1).number_format, iss, "Date")
        if d is None:
            iss.append(issue("error", "missing_date", "Date is empty."))
        _common_date_checks(d, None, today, iss)
        kind = {"installment": "receipt", "expense": "expense", "bill": "bill"}.get(typ)
        if kind is None:
            kind = next(iter(amounts), None)
            kind = {"installment": "receipt", "expense": "expense", "bill": "bill"}.get(kind)
            iss.append(issue("warning", "type_guessed", f"Type '{get('type') or ''}' not recognised; guessed {kind} from the amount column."))
        if len(amounts) > 1:
            iss.append(issue("error", "several_amounts", "More than one amount column is filled in on this row."))
        try:
            amt = _money(next(iter(amounts.values()), None))
        except (InvalidOperation, ValueError):
            amt = None
            iss.append(issue("error", "bad_amount", "Amount is not valid."))
        if amt is None or amt <= 0:
            iss.append(issue("error", "bad_amount", "Amount is missing or not greater than zero."))
        desc = str(get("description") or "").strip()
        ref = str(get("ref") or "").strip() or None
        mapped = {"kind": kind, "entry_date": d and str(d), "amount": amt and str(amt),
                  "description": desc or f"Imported row {r}", "notes": (str(get("notes")).strip() or None) if get("notes") else None,
                  "reference_no": ref, "party_text": (str(get("party")).strip() if get("party") else None),
                  "party_id": None, "money_account_id": None, "payment_method": guess_method(str(get("mode") or "")),
                  "source_cell": f"row {r}"}
        if kind == "receipt":
            mapped["receipt_kind"] = "customer_payment"
        elif kind == "expense":
            mapped.update(category_text=desc, category_code=guess_category(desc), expense_account_id=None)
        elif kind == "bill":
            mapped.update(direction="receivable", bill_no=ref, account_id=None)
            if not ref:
                iss.append(issue("error", "missing_bill_no", "A bill needs an invoice number."))
        out.append(ParsedRow(ws_f.title, r, cells, mapped, iss))
    return out
