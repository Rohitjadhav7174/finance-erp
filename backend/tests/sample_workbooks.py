"""SAMPLE workbooks for import tests. They copy the LAYOUT and the tricky features of the
owner's workbook (future date, date stored as a number, cell comments, a year-end total
that leaves rows out, an empty ledger template sheet) — but every value is invented."""
import io
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment

HEADER = ["Period", "Date ", "Opening Balance (₹)", "Payments", "Monthly Bill (₹)", "Other Expense / Charge (₹)",
          "Installment Received (₹)", "Expenses (₹)", "Cash Balance (₹)", "Outstanding Bill (₹)", "Period From",
          "Period To", "Period Shortfall (₹)", "Status"]
LEDGER_HEADER = ["Date", "Month", "Type", "Description / Particulars", "Bill Raised (₹)", "Installment Received (₹)",
                 "Expense Paid (₹)", "Payment Mode", "Invoice / Receipt No.", "Notes"]


def _with_cached_values(data: bytes, cached: dict[str, str]) -> bytes:
    """Writes Excel-style cached results into formula cells (as if Excel had saved the file)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = content.decode()
                for cell, value in cached.items():
                    xml = re.sub(rf'(<c r="{cell}"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/>', rf"\1<v>{value}</v>", xml)
                content = xml.encode()
            dst.writestr(item, content)
    return out.getvalue()


def balance_sheet(variant: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SAMPLE Site "
    ws["A1"] = "SAMPLE Trading Company – Monthly Balance Sheet"
    ws.append(HEADER)
    ws["A3"], ws["B3"], ws["C3"] = datetime(2025, 8, 1), datetime(2025, 8, 26), 11000
    ws["D3"], ws["E3"] = "Godown Bill", 3500
    ws["E3"].comment = Comment("======\nID#SAMPLE\nUnknown Author\nNotebook: SAMPLE bill figure not in the sheet.", "x")
    ws["D4"], ws["E4"] = "Diesel Payment", 9000
    ws["A6"], ws["B6"], ws["C6"] = datetime(2025, 9, 1), datetime(2025, 9, 4), 22000 + variant
    ws["A7"], ws["B7"], ws["C7"] = datetime(2025, 12, 1), datetime(2099, 12, 14), 33000   # future date (like row 12)
    ws["A8"] = "Total "
    ws["C8"] = "=SUM(C3:C7)"
    ws["A10"], ws["B10"], ws["C10"] = 46174, 46181, 44000                                # dates stored as numbers
    ws["D11"] = "Godown Bill"                                                            # category without amount
    ws["A12"] = "TOTAL / YEAR-END"
    ws["C12"] = "=C10"
    ws["E12"] = "=SUM(E5:E10)"                                                           # leaves out E3:E4
    ws["A15"], ws["C15"] = "Company", "SAMPLE Trading Company"
    ledger = wb.create_sheet("SAMPLE Ledger ")
    ledger["A1"] = "SAMPLE Trading Company – Transaction Ledger"
    ledger.append(LEDGER_HEADER)
    for r in range(3, 8):
        ledger[f"B{r}"] = f'=IF(A{r}="","",TEXT(A{r},"mmmm"))'
    b = io.BytesIO()
    wb.save(b)
    return _with_cached_values(b.getvalue(), {"C8": str(66000 + variant), "C12": "44000", "E12": "0"})


def ledger_workbook(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["SAMPLE ledger"])
    ws.append(["Date", "Type", "Description / Particulars", "Party", "Bill Raised (₹)", "Installment Received (₹)",
               "Expense Paid (₹)", "Payment Mode", "Invoice / Receipt No.", "Notes"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
