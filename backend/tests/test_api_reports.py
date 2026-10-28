"""Reports & dashboard figures checked against a hand-calculated SAMPLE scenario.

Scenario (company "RPT", all amounts SAMPLE):
  2025-04-01 opening cash 1,98,000
  2025-08-01 customer bill 99,000 (period Aug, due 31-08, installments 15-08 & 31-08 of 49,500)
  2025-08-05 supplier bill 5,000 (Other expenses)
  2025-08-20 received 49,500 cash against the bill           (installment 1)
  2025-08-22 advance 10,000 into bank (no receipt no.)       (liability, not income)
  2025-08-26 diesel 90,000 cash; godown 35,000 cash
  2025-08-27 transfer 20,000 cash -> bank                     (not money in/out)
  2025-08-28 paid supplier 5,000 from bank against its bill  (not a second expense)
  2025-09-04 received 49,500 bank against the bill           (bill now paid)
  2025-09-10 worker payment 75,000 cash, later reversed (reversal dated today)

Hand-calculated August: opening 1,98,000 · money in 59,500 (49,500 from customers) ·
money out 1,30,000 · income 99,000 · expenses 1,30,000 · closing cash+bank 1,27,500
(cash 1,02,500 + bank 25,000) · receivables 49,500 · advances 10,000 · shortfall 49,500.
"""
import csv
import io
import uuid

import pytest
from openpyxl import load_workbook

S = {}


@pytest.fixture(scope="module")
def rpt(seed):
    """Builds the scenario once, through the API, as the (SAMPLE) super admin."""
    from fastapi.testclient import TestClient

    from app import supabase_admin
    from app.main import app
    from conftest import FakeAuthAdmin, make_token

    app.dependency_overrides[supabase_admin.get_auth_admin] = lambda: FakeAuthAdmin()
    h = {"Authorization": f"Bearer {make_token(seed['users']['super'])}"}
    with TestClient(app) as c:
        code = "R" + uuid.uuid4().hex[:4].upper().translate(str.maketrans("0123456789", "ABCDEFGHIJ"))
        cid = c.post("/api/companies", json={"code": code, "name": "SAMPLE Report Co"}, headers=h).json()["id"]
        acc = {a["code"]: a["id"] for a in c.get(f"/api/companies/{cid}/accounts", headers=h).json()}
        bank = c.post(f"/api/companies/{cid}/accounts", headers=h,
                      json={"code": "BOB", "name": "SAMPLE BoB", "kind": "bank"}).json()["id"]

        def party(name, kind):
            return c.post(f"/api/companies/{cid}/parties", json={"name": name, "kind": kind}, headers=h).json()["id"]

        def post(kind, body):
            r = c.post(f"/api/companies/{cid}/entries/{kind}", json={**body, "action": "post"}, headers=h)
            assert r.status_code == 201, r.text
            return r.json()

        cust = party("=HYPERLINK(\"http://evil\")", "customer")  # formula-injection probe in a party name
        supp = party("SAMPLE Godown Owner", "supplier")
        post("opening", {"entry_date": "2025-04-01", "description": "SAMPLE opening",
                         "balances": [{"account_id": acc["CASH"], "side": "debit", "amount": "198000"}]})
        bill = post("bill", {"direction": "receivable", "bill_no": "INV-1", "party_id": cust, "entry_date": "2025-08-01",
                             "due_date": "2025-08-31", "period_from": "2025-08-01", "period_to": "2025-08-31",
                             "amount": "99000", "account_id": acc["INC-SALE"], "description": "SAMPLE Aug bill",
                             "installments": [{"due_date": "2025-08-15", "expected_amount": "49500"},
                                              {"due_date": "2025-08-31", "expected_amount": "49500"}]})["bill"]["id"]
        sbill = post("bill", {"direction": "payable", "bill_no": "SB-1", "party_id": supp, "entry_date": "2025-08-05",
                              "amount": "5000", "account_id": acc["EXP-OTH"], "description": "SAMPLE supplier bill"})["bill"]["id"]
        rc = lambda amt, d, money, ref: {"receipt_kind": "customer_payment", "party_id": cust, "entry_date": d,  # noqa: E731
                                          "amount": amt, "money_account_id": money, "payment_method": "cash",
                                          "reference_no": ref, "description": "SAMPLE installment",
                                          "allocations": [{"bill_id": bill, "amount": amt}]}
        post("receipt", rc("49500", "2025-08-20", acc["CASH"], "R-1"))
        post("receipt", {"receipt_kind": "advance", "party_id": cust, "entry_date": "2025-08-22", "amount": "10000",
                         "money_account_id": bank, "payment_method": "upi", "description": "SAMPLE advance"})
        for cat, amt in (("EXP-DSL", "90000"), ("EXP-GDN", "35000")):
            post("expense", {"entry_date": "2025-08-26", "expense_account_id": acc[cat], "amount": amt,
                             "money_account_id": acc["CASH"], "payment_method": "cash", "description": f"SAMPLE {cat}"})
        post("transfer", {"entry_date": "2025-08-27", "from_account_id": acc["CASH"], "to_account_id": bank,
                          "amount": "20000", "description": "SAMPLE deposit"})
        post("payment", {"party_id": supp, "entry_date": "2025-08-28", "amount": "5000", "money_account_id": bank,
                         "payment_method": "bank_transfer", "reference_no": "UTR-1", "description": "SAMPLE pay supplier",
                         "allocations": [{"bill_id": sbill, "amount": "5000"}]})
        post("receipt", {**rc("49500", "2025-09-04", bank, "R-2"), "payment_method": "upi"})
        wrk = post("expense", {"entry_date": "2025-09-10", "expense_account_id": acc["EXP-WRK"], "amount": "75000",
                               "money_account_id": acc["CASH"], "payment_method": "cash", "description": "SAMPLE workers"})
        c.post(f"/api/entries/{wrk['id']}/reversal-request", json={"reason": "SAMPLE duplicate"}, headers=h)
        appr = c.get(f"/api/companies/{cid}/approvals", headers=h).json()[0]["id"]
        assert c.post(f"/api/approvals/{appr}/decide", json={"decision": "approve", "note": "owner checked"},
                      headers=h).status_code == 200
        assert c.post(f"/api/entries/{wrk['id']}/reverse", headers=h).status_code == 201
    app.dependency_overrides.clear()
    return {"cid": cid, "bill": bill, "acc": acc, "bank": bank}


def get(client, h, rpt, path, **params):
    r = client.get(f"/api/companies/{rpt['cid']}/{path}", params=params, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_dashboard_august_figures(client, auth, rpt):
    f = get(client, auth("super"), rpt, "dashboard", date_from="2025-08-01", date_to="2025-08-31")["figures"]
    assert f["opening_cash"] == "198000.00"
    assert f["money_in"] == "59500.00" and f["collected_from_customers"] == "49500.00"
    assert f["money_out"] == "130000.00"
    assert f["income"] == "99000.00" and f["expenses"] == "130000.00" and f["surplus"] == "-31000.00"
    assert (f["closing_cash"], f["cash"], f["bank"]) == ("127500.00", "102500.00", "25000.00")
    assert (f["receivables"], f["payables"], f["advances"]) == ("49500.00", "0.00", "10000.00")
    assert f["bills_raised"] == "99000.00"
    assert f["other_movement"] == "0.00"


def test_cash_movement_always_reconciles(client, auth, rpt):
    """opening + money in - money out + opening balances/adjustments = closing, for any period."""
    from decimal import Decimal as D
    for d_from, d_to in (("2025-04-01", "2025-08-31"), ("2025-04-01", "2026-12-31"), ("2025-08-15", "2025-09-15")):
        f = get(client, auth("super"), rpt, "dashboard", date_from=d_from, date_to=d_to)["figures"]
        assert D(f["opening_cash"]) + D(f["money_in"]) - D(f["money_out"]) + D(f["other_movement"]) == D(f["closing_cash"]), (d_from, f)
    f = get(client, auth("super"), rpt, "dashboard", date_from="2025-04-01", date_to="2025-08-31")["figures"]
    assert f["other_movement"] == "198000.00"  # the opening balance entry


def test_dashboard_september_and_reversal(client, auth, rpt):
    f = get(client, auth("super"), rpt, "dashboard", date_from="2025-09-01", date_to="2025-09-30")["figures"]
    assert f["receivables"] == "0.00" and f["money_in"] == "49500.00"
    assert f["expenses"] == "75000.00"  # the reversal is dated today, so it lands in the current period
    whole = get(client, auth("super"), rpt, "reports/income_expense", date_from="2025-04-01")
    assert not any(r["account"] == "Worker Payment" and r["amount"] != "0.00" for r in whole["rows"])


def test_monthly_balance_sheet(client, auth, rpt):
    rep = get(client, auth("super"), rpt, "reports/monthly_balance_sheet", date_from="2025-08-01", date_to="2025-09-30")
    aug, sep = rep["rows"]
    assert (aug["period"], aug["opening"], aug["bills_raised"], aug["installments_received"]) == \
        ("Aug-2025", "198000.00", "99000.00", "49500.00")
    assert (aug["other_receipts"], aug["closing"], aug["outstanding_bill"], aug["shortfall"], aug["status"]) == \
        ("10000.00", "127500.00", "49500.00", "49500.00", "Shortfall")
    assert "Diesel 90000.00" in aug["payments_particulars"]
    assert (sep["opening"], sep["outstanding_bill"], sep["shortfall"]) == ("127500.00", "0.00", "0.00")
    assert any("awaiting owner confirmation" in n for n in rep["notes"])


def test_cash_and_bank_books(client, auth, rpt):
    cash = get(client, auth("super"), rpt, "reports/cash_book", date_from="2025-08-01", date_to="2025-08-31")
    assert cash["totals"]["opening"] == "198000.00" and cash["totals"]["balance"] == "102500.00"
    assert [r["balance"] for r in cash["rows"]][-1] == "102500.00"
    bank = get(client, auth("super"), rpt, "reports/bank_book", date_from="2025-08-01", date_to="2025-08-31")
    assert bank["totals"]["balance"] == "25000.00"
    assert (bank["totals"]["receipt"], bank["totals"]["payment"]) == ("30000.00", "5000.00")


def test_receivables_aging_and_installments(client, auth, rpt):
    rec = get(client, auth("super"), rpt, "reports/receivables", date_to="2025-08-31", date_from="2025-04-01")
    row = rec["rows"][0]
    assert (row["balance"], row["not_due"], row["on_account"]) == ("49500.00", "49500.00", "0.00")
    rec_later = get(client, auth("super"), rpt, "reports/receivables", date_to="2025-10-31", date_from="2025-04-01")
    assert rec_later["rows"] == []
    inst = get(client, auth("super"), rpt, "reports/installments", date_from="2025-08-01", date_to="2025-08-31")
    assert [(r["expected"], r["collected"], r["status"]) for r in inst["rows"]] == \
        [("49500.00", "49500.00", "paid"), ("49500.00", "49500.00", "paid")]


def test_registers_and_income_expense(client, auth, rpt):
    bills = get(client, auth("super"), rpt, "reports/bill_register", date_from="2025-08-01", date_to="2025-08-31")
    assert bills["totals"] == {"amount": "104000.00", "paid": "54500.00", "outstanding": "49500.00"}
    ie = get(client, auth("super"), rpt, "reports/income_expense", date_from="2025-08-01", date_to="2025-08-31")
    assert ie["totals"] == {"income": "99000.00", "expenses": "130000.00", "amount": "-31000.00"}
    tr = get(client, auth("super"), rpt, "reports/transaction_register", date_from="2025-08-01", date_to="2025-08-31")
    assert len(tr["rows"]) == 8  # 2 bills, 2 receipts, 2 expenses, transfer, supplier payment


def test_pending_work(client, auth, rpt):
    p = get(client, auth("super"), rpt, "pending")
    assert p["counts"].get("missing_reference") == 1  # the advance has no receipt number
    assert "overdue_bill" not in p["counts"]  # everything was paid
    assert "unreconciled_bank" not in p["counts"]  # no reconciliation started, so not flagged
    only = get(client, auth("super"), rpt, "pending", kind="missing_reference")
    assert len(only["items"]) == 1 and only["items"][0]["amount"] == "10000.00"


def test_exports_csv_and_xlsx_are_safe_and_exact(client, auth, rpt):
    h = auth("super")
    r = client.get(f"/api/companies/{rpt['cid']}/reports/receivables",
                   params={"date_to": "2025-08-31", "date_from": "2025-04-01", "format": "csv"}, headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    party_cells = [row[0] for row in rows if row and "HYPERLINK" in row[0]]
    assert party_cells and all(c.startswith("'=") for c in party_cells), party_cells  # never a live formula
    assert any("49500.00" in row for row in rows)

    r = client.get(f"/api/companies/{rpt['cid']}/reports/monthly_balance_sheet",
                   params={"date_from": "2025-08-01", "date_to": "2025-08-31", "format": "xlsx"}, headers=h)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert 127500 in values and 198000 in values  # real numbers, not text
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


def test_reports_respect_company_access(client, auth, rpt):
    for who in ("JGT:company_admin", "ART:viewer"):
        r = client.get(f"/api/companies/{rpt['cid']}/reports/cash_book", headers=auth(who))
        assert r.status_code == 404
        assert client.get(f"/api/companies/{rpt['cid']}/dashboard", headers=auth(who)).status_code == 404


def test_audit_report_admins_only(client, auth, seed):
    art = seed["companies"]["ART"]
    assert client.get(f"/api/companies/{art}/reports/audit", headers=auth("ART:accountant")).status_code == 403
    assert client.get(f"/api/companies/{art}/reports/audit", headers=auth("ART:company_admin")).status_code == 200


def test_consolidated_requires_permission(client, auth, rpt):
    assert client.get("/api/consolidated", headers=auth("ART:company_admin")).status_code == 403
    r = client.get("/api/consolidated", headers=auth("super"), params={"date_to": "2025-08-31"})
    assert r.status_code == 200
    rpt_row = next(c for c in r.json()["companies"] if c["company"]["id"] == rpt["cid"])
    assert rpt_row["figures"]["receivables"] == "49500.00"


def test_global_search(client, auth, rpt):
    h = auth("super")
    r = client.get("/api/search", params={"q": "90,000"}, headers=h).json()
    assert any(e["amount"] == "90000.00" for e in r["entries"])
    r = client.get("/api/search", params={"q": "inv-1"}, headers=h).json()
    assert any(b["bill_no"] == "INV-1" for b in r["bills"])
    # a user of another company finds nothing from this one
    r = client.get("/api/search", params={"q": "inv-1"}, headers=auth("JGT:viewer")).json()
    assert not any(b["bill_no"] == "INV-1" for b in r["bills"])


def test_pending_lists_scheduled_bills_by_installment_only(client, auth, seed):
    """A bill with installments is followed up installment by installment, never listed twice."""
    art = seed["companies"]["ART"]
    h = auth("ART:accountant")
    acc = {a["code"]: a["id"] for a in client.get(f"/api/companies/{art}/accounts", headers=h).json()}
    cust = client.post(f"/api/companies/{art}/parties", json={"name": f"SAMPLE dedupe {uuid.uuid4().hex[:5]}",
                                                               "kind": "customer"}, headers=h).json()["id"]
    r = client.post(f"/api/companies/{art}/entries/bill", headers=h, json={
        "direction": "receivable", "bill_no": f"DD-{uuid.uuid4().hex[:5]}", "party_id": cust, "entry_date": "2025-05-01",
        "due_date": "2025-06-30", "amount": "1000", "account_id": acc["INC-SALE"], "description": "SAMPLE", "action": "post",
        "installments": [{"due_date": "2025-05-31", "expected_amount": "500"}, {"due_date": "2025-06-30", "expected_amount": "500"}]})
    bill_id = r.json()["bill"]["id"]
    items = [i for i in client.get(f"/api/companies/{art}/pending", headers=h).json()["items"] if i["ref_id"] == bill_id]
    assert sorted(i["kind"] for i in items) == ["overdue_installment", "overdue_installment"]
